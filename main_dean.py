import sqlite3
import string
import random
from openpyxl import Workbook
import pandas as pd
from PyQt5.QtWidgets import QApplication, QMainWindow, QMenu, QMenuBar, QAction, QFileDialog, QTableWidget, \
    QTableWidgetItem, QVBoxLayout, QWidget, QHeaderView, QDialog, QLabel, QLineEdit, QPushButton, QComboBox, QHBoxLayout, QMessageBox
import sys
from transliterate import translit
import csv
import configparser
import os
import re

# Load database path from config.ini with error handling
config = configparser.ConfigParser()
DB_PATH = None
if os.path.exists('config.ini'):
    config.read('config.ini')
    try:
        DB_PATH = config['Database']['Path']
    except (KeyError, configparser.NoSectionError):
        print("Error: 'Database' section or 'Path' key not found in config.ini. Please check the file format.")

class MainWindow(QMainWindow):
    def __init__(self):
        super().__init__()
        self.current_db_path = DB_PATH
        self.current_account_id = None
        self.current_user_role = None
        self.current_user_faculty_access = []
        self.initUI()

    def initUI(self):
        print("Initializing UI")
        menubar = self.menuBar()






        delete_action = QAction('Удалить пользователя', self)
        delete_action.triggered.connect(self.delete_user)
        add_action = QAction('Добавить пользователя', self)
        add_action.triggered.connect(self.add_user_form)
        edit_action = QAction('Редактировать пользователя', self)
        edit_action.triggered.connect(self.edit_user_dialog)


        add_faculty_action = QAction('Добавить факультет', self)
        add_faculty_action.triggered.connect(self.add_faculty_form)
        add_division_action = QAction('Добавить группу', self)
        add_division_action.triggered.connect(self.add_division_form)
        manage_accounts_action = QAction('Управление учётными записями', self)
        manage_accounts_action.triggered.connect(self.manage_accounts_form)


        load_export_menu = QMenu('Загрузка/Выгрузка', self)
        export_to_excel_action = QAction('Выгрузить в Excel', self)
        export_to_excel_action.triggered.connect(self.export_selected_to_excel)
        load_export_menu.addAction(export_to_excel_action)
        export_to_CSV_action = QAction('Выгрузить в CSV', self)
        export_to_CSV_action.triggered.connect(self.export_selected_to_CSV)
        load_export_menu.addAction(export_to_CSV_action)

        menubar.addMenu(load_export_menu)

        self.setWindowTitle('Управление пользователями')
        self.setGeometry(100, 100, 920, 600)



    def show_login_dialog(self):
        print("Showing login dialog")
        login_dialog = QDialog(self)
        login_dialog.setWindowTitle("Авторизация")

        login_label = QLabel("Логин:")
        login_edit = QLineEdit()
        password_label = QLabel("Пароль:")
        password_edit = QLineEdit()
        password_edit.setEchoMode(QLineEdit.Password)
        login_button = QPushButton("Войти")
        cancel_button = QPushButton("Отмена")

        layout = QVBoxLayout()
        layout.addWidget(login_label)
        layout.addWidget(login_edit)
        layout.addWidget(password_label)
        layout.addWidget(password_edit)
        layout.addWidget(login_button)
        layout.addWidget(cancel_button)
        login_dialog.setLayout(layout)

        # Flag to track authentication success
        self.auth_success = False

        def on_login():
            self.authenticate(login_edit.text(), password_edit.text(), login_dialog)

        def on_cancel():
            print("Login dialog canceled")
            login_dialog.reject()

        login_button.clicked.connect(on_login)
        cancel_button.clicked.connect(on_cancel)
        login_dialog.rejected.connect(lambda: print("Login dialog closed"))

        # Keep dialog open until successful login or cancel/close
        while not self.auth_success:
            if login_dialog.exec_() != QDialog.Accepted:
                print("Exiting due to dialog rejection or close")
                return False  # Signal to exit program
            if self.auth_success:
                break

        return True  # Signal successful authentication

    def authenticate(self, login, password, dialog):
        print(f"Authenticating user: {login}")
        try:
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT id, role, faculty_access FROM Accounts WHERE login=? AND password=?", (login, password))
            account = cursor.fetchone()
            if account:
                print(f"Account found: ID={account[0]}, Role={account[1]}, Faculty_access={account[2]}")
                self.current_account_id = account[0]
                self.current_user_role = account[1]
                self.current_user_faculty_access = account[2].split(',') if account[2] and isinstance(account[2], str) else []
                self.auth_success = True
                conn.close()
                dialog.accept()  # Close dialog and proceed
            else:
                conn.close()
                print("Authentication failed: Invalid login or password")
                QMessageBox.warning(self, "Ошибка", "Не успешная авторизация")
                # Do not reject dialog; let loop retry
        except Exception as e:
            print(f"Authentication error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка авторизации: {e}")
            conn.close()

    def load_data(self):
        print(f"Loading data from {self.current_db_path}")
        try:
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            query = "SELECT id, surname, name, patronymic, login, password, division, post, faculty FROM user"
            params = []

            if self.current_user_role != 'admin':
                print(f"Applying role-based filter for role: {self.current_user_role}")
                if self.current_user_role == 'dean':
                    cursor.execute("SELECT faculty_access FROM Accounts WHERE id=?", (self.current_account_id,))
                    dean_faculty = cursor.fetchone()
                    print(f"Dean faculty access: {dean_faculty}")
                    if dean_faculty and dean_faculty[0]:
                        query += " WHERE faculty = ?"
                        params.append(dean_faculty[0])
                    else:
                        print("No faculty access for dean, returning empty result")
                        query += " WHERE 1=0"
                else:  # user role
                    print(f"User faculty access: {self.current_user_faculty_access}")
                    if self.current_user_faculty_access:
                        placeholders = ','.join('?' for _ in self.current_user_faculty_access)
                        query += f" WHERE faculty IN ({placeholders})"
                        params.extend(self.current_user_faculty_access)
                    else:
                        print("No faculty access for user, returning empty result")
                        query += " WHERE 1=0"

            print(f"Executing query: {query} with params: {params}")
            cursor.execute(query, params)
            data = cursor.fetchall()
            print(f"Fetched {len(data)} rows: {data}")

            table_area = QWidget()
            layout = QVBoxLayout(table_area)

            filter_layout = QHBoxLayout()
            self.last_name_filter = QLineEdit()
            self.first_name_filter = QLineEdit()
            self.middle_name_filter = QLineEdit()
            self.subdivision_filter = QComboBox()
            self.position_filter = QComboBox()
            self.faculty_filter = QComboBox()

            subdivisions = self.get_divisions()
            positions = self.get_posts()
            faculties = self.get_faculties()
            print(f"Subdivisions: {subdivisions}, Positions: {positions}, Faculties: {faculties}")
            self.subdivision_filter.addItem("Все")
            self.subdivision_filter.addItems(subdivisions)
            self.position_filter.addItem("Все")
            self.position_filter.addItems(positions)
            self.faculty_filter.addItem("Все")
            self.faculty_filter.addItems(faculties)

            filter_layout.addWidget(QLabel("Фамилия:"))
            filter_layout.addWidget(self.last_name_filter)
            filter_layout.addWidget(QLabel("Имя:"))
            filter_layout.addWidget(self.first_name_filter)
            filter_layout.addWidget(QLabel("Отчество:"))
            filter_layout.addWidget(self.middle_name_filter)
            filter_layout.addWidget(QLabel("Подразделение:"))
            filter_layout.addWidget(self.subdivision_filter)
            filter_layout.addWidget(QLabel("Должность:"))
            filter_layout.addWidget(self.position_filter)
            filter_layout.addWidget(QLabel("Факультет:"))
            filter_layout.addWidget(self.faculty_filter)

            filter_button = QPushButton("Применить")
            filter_button.clicked.connect(self.apply_filters)
            filter_layout.addWidget(filter_button)

            layout.addLayout(filter_layout)

            self.table = QTableWidget()
            self.table.setColumnCount(9)
            self.table.setRowCount(len(data))
            self.table.setHorizontalHeaderLabels(
                ["ID", "Фамилия", "Имя", "Отчество", "Логин", "Пароль", "Подразделение", "Должность", "Факультет"])

            for row, row_data in enumerate(data):
                for col, item in enumerate(row_data):
                    item_text = str(item) if item is not None else ''
                    self.table.setItem(row, col, QTableWidgetItem(item_text))

            self.table.hideColumn(0)
            self.table.setSelectionBehavior(QTableWidget.SelectRows)

            header = self.table.horizontalHeader()
            header.setSectionResizeMode(QHeaderView.Stretch)
            header.sectionClicked.connect(self.sort_table)

            layout.addWidget(self.table)
            self.setCentralWidget(table_area)
            conn.close()
            print("Data loaded successfully")
        except Exception as e:
            print(f"Load data error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка загрузки данных: {e}")
            conn.close()

    def apply_filters(self):
        print("Applying filters")
        try:
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            query = "SELECT id, surname, name, patronymic, login, password, division, post, faculty FROM user WHERE 1=1"
            params = []

            if self.current_user_role != 'admin':
                if self.current_user_role == 'dean':
                    cursor.execute("SELECT faculty_access FROM Accounts WHERE id=?", (self.current_account_id,))
                    dean_faculty = cursor.fetchone()
                    if dean_faculty and dean_faculty[0]:
                        query += " AND faculty = ?"
                        params.append(dean_faculty[0])
                    else:
                        query += " AND 1=0"
                else:
                    if self.current_user_faculty_access:
                        placeholders = ','.join('?' for _ in self.current_user_faculty_access)
                        query += f" AND faculty IN ({placeholders})"
                        params.extend(self.current_user_faculty_access)
                    else:
                        query += " AND 1=0"

            if self.last_name_filter.text():
                query += " AND surname LIKE ?"
                params.append(f"%{self.last_name_filter.text()}%")
            if self.first_name_filter.text():
                query += " AND name LIKE ?"
                params.append(f"%{self.first_name_filter.text()}%")
            if self.middle_name_filter.text():
                query += " AND patronymic LIKE ?"
                params.append(f"%{self.middle_name_filter.text()}%")
            if self.subdivision_filter.currentText() != "Все":
                query += " AND division = ?"
                params.append(self.subdivision_filter.currentText())
            if self.position_filter.currentText() != "Все":
                query += " AND post = ?"
                params.append(self.position_filter.currentText())
            if self.faculty_filter.currentText() != "Все":
                query += " AND faculty = ?"
                params.append(self.faculty_filter.currentText())

            print(f"Filter query: {query} with params: {params}")
            cursor.execute(query, params)
            data = cursor.fetchall()

            self.table.setRowCount(len(data))
            for row, row_data in enumerate(data):
                for col, item in enumerate(row_data):
                    item_text = str(item) if item is not None else ''
                    self.table.setItem(row, col, QTableWidgetItem(item_text))

            self.table.hideColumn(0)
            conn.close()
        except Exception as e:
            print(f"Filter error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка применения фильтров: {e}")
            conn.close()

    def sort_table(self, column):
        print(f"Sorting table by column {column}")
        try:
            sort_order = self.table.horizontalHeader().sortIndicatorOrder()
            self.table.sortItems(column, sort_order)
        except Exception as e:
            print(f"Sort error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка сортировки: {e}")

    def delete_user(self):
        print("Deleting user")
        selected_rows = self.table.selectedItems()
        if not selected_rows:
            QMessageBox.warning(self, "Ошибка", "Выберите пользователя для удаления")
            return
        ids_to_delete = set()
        for item in selected_rows:
            row = item.row()
            id_item = self.table.item(row, 0)
            if id_item:
                ids_to_delete.add(int(id_item.text()))

        try:
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            for user_id in ids_to_delete:
                cursor.execute("DELETE FROM user WHERE id=?", (user_id,))
            conn.commit()
            conn.close()
            self.load_data()
        except Exception as e:
            print(f"Delete user error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка удаления пользователя: {e}")
            conn.close()

    def add_user_form(self):
        print("Opening add user form")
        add_dialog = QDialog(self)
        add_dialog.setWindowTitle("Добавить пользователя")

        last_name_label = QLabel("Фамилия:")
        last_name_edit = QLineEdit()
        first_name_label = QLabel("Имя:")
        first_name_edit = QLineEdit()
        middle_name_label = QLabel("Отчество:")
        middle_name_edit = QLineEdit()
        login_label = QLabel("Логин:")
        login_edit = QLineEdit()
        password_label = QLabel("Пароль:")
        password_edit = QLineEdit()
        subdivision_label = QLabel("Подразделение:")
        subdivision_edit = QComboBox()
        position_label = QLabel("Должность:")
        position_edit = QComboBox()
        faculty_label = QLabel("Факультет:")
        faculty_edit = QComboBox()

        add_button = QPushButton("Добавить")
        cancel_button = QPushButton("Отмена")
        generate_login_button = QPushButton("Сгенерировать логин")
        generate_password_button = QPushButton("Сгенерировать пароль")

        subdivisions = self.get_divisions()
        positions = self.get_posts()
        faculties = self.get_faculties()
        subdivision_edit.addItems(subdivisions)
        position_edit.addItems(positions)
        faculty_edit.addItems(faculties)

        layout = QVBoxLayout()
        layout.addWidget(last_name_label)
        layout.addWidget(last_name_edit)
        layout.addWidget(first_name_label)
        layout.addWidget(first_name_edit)
        layout.addWidget(middle_name_label)
        layout.addWidget(middle_name_edit)
        layout.addWidget(login_label)
        layout.addWidget(login_edit)
        layout.addWidget(generate_login_button)
        layout.addWidget(password_label)
        layout.addWidget(password_edit)
        layout.addWidget(generate_password_button)
        layout.addWidget(subdivision_label)
        layout.addWidget(subdivision_edit)
        layout.addWidget(position_label)
        layout.addWidget(position_edit)
        layout.addWidget(faculty_label)
        layout.addWidget(faculty_edit)
        layout.addWidget(add_button)
        layout.addWidget(cancel_button)

        add_dialog.setLayout(layout)

        generate_login_button.clicked.connect(lambda: self.generate_login(last_name_edit, first_name_edit, middle_name_edit, login_edit))
        generate_password_button.clicked.connect(lambda: self.generate_password(password_edit))

        add_button.clicked.connect(lambda: self.add_user(
            last_name_edit.text(),
            first_name_edit.text(),
            middle_name_edit.text(),
            login_edit.text(),
            password_edit.text(),
            subdivision_edit.currentText(),
            position_edit.currentText(),
            faculty_edit.currentText(),
            add_dialog
        ))

        cancel_button.clicked.connect(add_dialog.close)

        add_dialog.exec_()

    def generate_login(self, last_name_edit, first_name_edit, middle_name_edit, login_edit):
        print("Generating login")
        try:
            last_name = last_name_edit.text().strip()
            first_name = first_name_edit.text().strip()
            middle_name = middle_name_edit.text().strip()

            last_name_translit = translit(last_name, 'ru', reversed=True).capitalize()
            first_initial = translit(first_name[0], 'ru', reversed=True).upper() if first_name else ''
            middle_initial = translit(middle_name[0], 'ru', reversed=True).upper() if middle_name else ''

            # Combine and clean the login
            base_login = f"{last_name_translit}{first_initial}{middle_initial}"
            # Remove single quotes and other unwanted characters (e.g., non-alphanumeric except underscore)
            clean_login = re.sub(r"['\W]+", '', base_login)  # Removes single quotes and non-alphanumeric characters
            login = clean_login

            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT login FROM user WHERE login=?", (login,))
            existing_login = cursor.fetchone()

            counter = 1
            while existing_login:
                login = f"{clean_login}{counter}"
                cursor.execute("SELECT login FROM user WHERE login=?", (login,))
                existing_login = cursor.fetchone()
                counter += 1

            conn.close()
            login_edit.setText(login)
        except Exception as e:
            print(f"Generate login error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка генерации логина: {e}")

    def generate_password(self, password_edit):
        print("Generating password")
        try:
            allowed_chars = 'ABCDEFGHIJKMNPQRSTUVWXYZabcdefghijkmnpqrstuvwxyz23456789'
            password = ''.join(random.choice(allowed_chars) for _ in range(8))
            password_edit.setText(password)
        except Exception as e:
            print(f"Generate password error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка генерации пароля: {e}")

    def generate_account_login(self, login_edit):
        print("Generating account login")
        try:
            base_login = ''.join(random.choice(string.ascii_letters) for _ in range(6))
            login = base_login

            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT login FROM Accounts WHERE login=?", (login,))
            existing_login = cursor.fetchone()

            counter = 1
            while existing_login:
                login = f"{base_login}{counter}"
                cursor.execute("SELECT login FROM Accounts WHERE login=?", (login,))
                existing_login = cursor.fetchone()
                counter += 1

            conn.close()
            login_edit.setText(login)
        except Exception as e:
            print(f"Generate account login error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка генерации логина учётной записи: {e}")

    def manage_accounts_form(self):
        print("Opening manage accounts form")
        dialog = QDialog(self)
        dialog.setWindowTitle("Управление учётными записями")

        login_label = QLabel("Логин:")
        login_edit = QLineEdit()
        password_label = QLabel("Пароль:")
        password_edit = QLineEdit()
        role_label = QLabel("Роль:")
        role_edit = QComboBox()
        role_edit.addItems(['admin', 'dean', 'user'])
        faculty_access_label = QLabel("Доступ к факультетам")
        faculty_access_edit = QLineEdit()

        add_button = QPushButton("Добавить")
        cancel_button = QPushButton("Отмена")
        generate_login_button = QPushButton("Сгенерировать логин")
        generate_password_button = QPushButton("Сгенерировать пароль")

        layout = QVBoxLayout()
        layout.addWidget(login_label)
        layout.addWidget(login_edit)
        layout.addWidget(generate_login_button)
        layout.addWidget(password_label)
        layout.addWidget(password_edit)
        layout.addWidget(generate_password_button)
        layout.addWidget(role_label)
        layout.addWidget(role_edit)
        layout.addWidget(faculty_access_label)
        layout.addWidget(faculty_access_edit)
        layout.addWidget(add_button)
        layout.addWidget(cancel_button)

        dialog.setLayout(layout)

        generate_login_button.clicked.connect(lambda: self.generate_account_login(login_edit))
        generate_password_button.clicked.connect(lambda: self.generate_password(password_edit))

        add_button.clicked.connect(lambda: self.add_account(
            login_edit.text(),
            password_edit.text(),
            role_edit.currentText(),
            faculty_access_edit.text(),
            dialog
        ))

        cancel_button.clicked.connect(dialog.close)

        dialog.exec_()

    def add_account(self, login, password, role, faculty_access, dialog):
        print(f"Adding account: {login}")
        try:
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            cursor.execute("""
                INSERT INTO Accounts (login, password, role, faculty_access)
                VALUES (?, ?, ?, ?)
            """, (login, password, role, faculty_access or None))
            conn.commit()
            conn.close()
            dialog.close()
            QMessageBox.information(self, "Успех", "Учётная запись добавлена")
        except sqlite3.IntegrityError as e:
            print(f"Add account integrity error: {e}")
            QMessageBox.warning(self, "Ошибка", f"Не удалось добавить учётную запись: {e}")
            conn.close()
        except Exception as e:
            print(f"Add account error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка добавления учётной записи: {e}")
            conn.close()

    def add_user(self, last_name, first_name, middle_name, login, password, subdivision, position, faculty, dialog):
        print("Adding user")
        try:
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            cursor.execute("""
                INSERT INTO user (surname, name, patronymic, login, password, division, post, faculty)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
            """, (last_name, first_name, middle_name or None, login, password, subdivision or None, position or None, faculty or None))
            conn.commit()
            conn.close()
            self.load_data()
            dialog.close()
        except sqlite3.IntegrityError as e:
            print(f"Add user integrity error: {e}")
            QMessageBox.warning(self, "Ошибка", f"Не удалось добавить пользователя: {e}")
            conn.close()
        except Exception as e:
            print(f"Add user error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка добавления пользователя: {e}")
            conn.close()

    def edit_user_dialog(self):
        print("Opening edit user dialog")
        selected_rows = self.table.selectedItems()
        if not selected_rows:
            QMessageBox.warning(self, "Ошибка", "Выберите пользователя для редактирования")
            return

        row = selected_rows[0].row()
        id_item = self.table.item(row, 0)

        edit_dialog = QDialog(self)
        edit_dialog.setWindowTitle("Редактировать пользователя")

        last_name_label = QLabel("Фамилия:")
        last_name_edit = QLineEdit(self.table.item(row, 1).text())
        first_name_label = QLabel("Имя:")
        first_name_edit = QLineEdit(self.table.item(row, 2).text())
        middle_name_label = QLabel("Отчество:")
        middle_name_edit = QLineEdit(self.table.item(row, 3).text() or '')
        login_label = QLabel("Логин:")
        login_edit = QLineEdit(self.table.item(row, 4).text())
        password_label = QLabel("Пароль:")
        password_edit = QLineEdit(self.table.item(row, 5).text())
        subdivision_label = QLabel("Подразделение:")
        subdivision_edit = QComboBox()
        position_label = QLabel("Должность:")
        position_edit = QComboBox()
        faculty_label = QLabel("Факультет:")
        faculty_edit = QComboBox()

        subdivisions = self.get_divisions()
        positions = self.get_posts()
        faculties = self.get_faculties()
        subdivision_edit.addItems(subdivisions)
        position_edit.addItems(positions)
        faculty_edit.addItems(faculties)

        subdivision_edit.setCurrentText(self.table.item(row, 6).text() if self.table.item(row, 6) else '')
        position_edit.setCurrentText(self.table.item(row, 7).text() if self.table.item(row, 7) else '')
        faculty_edit.setCurrentText(self.table.item(row, 8).text() if self.table.item(row, 8) else '')

        edit_button = QPushButton("Сохранить")
        cancel_button = QPushButton("Отмена")
        generate_login_button = QPushButton("Сгенерировать логин")
        generate_password_button = QPushButton("Сгенерировать пароль")

        layout = QVBoxLayout()
        layout.addWidget(last_name_label)
        layout.addWidget(last_name_edit)
        layout.addWidget(first_name_label)
        layout.addWidget(first_name_edit)
        layout.addWidget(middle_name_label)
        layout.addWidget(middle_name_edit)
        layout.addWidget(login_label)
        layout.addWidget(login_edit)
        layout.addWidget(generate_login_button)
        layout.addWidget(password_label)
        layout.addWidget(password_edit)
        layout.addWidget(generate_password_button)
        layout.addWidget(subdivision_label)
        layout.addWidget(subdivision_edit)
        layout.addWidget(position_label)
        layout.addWidget(position_edit)
        layout.addWidget(faculty_label)
        layout.addWidget(faculty_edit)
        layout.addWidget(edit_button)
        layout.addWidget(cancel_button)

        edit_dialog.setLayout(layout)

        generate_login_button.clicked.connect(lambda: self.generate_login(last_name_edit, first_name_edit, middle_name_edit, login_edit))
        generate_password_button.clicked.connect(lambda: self.generate_password(password_edit))

        edit_button.clicked.connect(lambda: self.edit_user(
            id_item.text(),
            last_name_edit.text(),
            first_name_edit.text(),
            middle_name_edit.text(),
            login_edit.text(),
            password_edit.text(),
            subdivision_edit.currentText(),
            position_edit.currentText(),
            faculty_edit.currentText(),
            edit_dialog
        ))

        cancel_button.clicked.connect(edit_dialog.close)

        edit_dialog.exec_()

    def edit_user(self, user_id, last_name, first_name, middle_name, login, password, subdivision, position, faculty, dialog):
        print(f"Editing user ID: {user_id}")
        try:
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            cursor.execute("""
                UPDATE user
                SET surname = ?, name = ?, patronymic = ?, login = ?, password = ?, division = ?, post = ?, faculty = ?
                WHERE id = ?
            """, (last_name, first_name, middle_name or None, login, password, subdivision or None, position or None, faculty or None, user_id))
            conn.commit()
            conn.close()
            self.load_data()
            dialog.close()
        except sqlite3.IntegrityError as e:
            print(f"Edit user integrity error: {e}")
            QMessageBox.warning(self, "Ошибка", f"Не удалось обновить пользователя: {e}")
            conn.close()
        except Exception as e:
            print(f"Edit user error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка редактирования пользователя: {e}")
            conn.close()

    def generate_login_import(self, last_name, first_name, middle_name):
        print("Generating login for import")
        try:
            last_name_translit = translit(str(last_name).strip(), 'ru', reversed=True).capitalize()
            first_initial = translit(str(first_name).strip()[0], 'ru', reversed=True).upper() if str(
                first_name).strip() else ''
            middle_initial = translit(str(middle_name).strip()[0], 'ru', reversed=True).upper() if str(
                middle_name).strip() else ''

            # Combine and clean the login
            base_login = f"{last_name_translit}{first_initial}{middle_initial}"
            clean_login = re.sub(r"['\W]+", '', base_login)  # Removes single quotes and non-alphanumeric characters
            login = clean_login

            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT login FROM user WHERE login=?", (login,))
            existing_login = cursor.fetchone()

            counter = 1
            while existing_login:
                login = f"{clean_login}{counter}"
                cursor.execute("SELECT login FROM user WHERE login=?", (login,))
                existing_login = cursor.fetchone()
                counter += 1

            conn.close()
            return login
        except Exception as e:
            print(f"Generate login import error: {e}")
            raise ValueError(f"Ошибка генерации логина для импорта: {e}")

    def generate_password_import(self):
        print("Generating password for import")
        try:
            allowed_chars = 'ABCDEFGHIJKMNPQRSTUVWXYZabcdefghijkmnpqrstuvwxyz23456789'
            password = ''.join(random.choice(allowed_chars) for _ in range(8))
            return password
        except Exception as e:
            print(f"Generate password import error: {e}")
            raise ValueError(f"Ошибка генерации пароля для импорта: {e}")

    def add_faculty_form(self):
        print("Opening add faculty form")
        add_dialog = QDialog(self)
        add_dialog.setWindowTitle("Добавить факультет")

        short_name_label = QLabel("Краткое название:")
        short_name_edit = QLineEdit()
        full_name_label = QLabel("Полное название:")
        full_name_edit = QLineEdit()

        add_button = QPushButton("Добавить")
        cancel_button = QPushButton("Отмена")

        layout = QVBoxLayout()
        layout.addWidget(short_name_label)
        layout.addWidget(short_name_edit)
        layout.addWidget(full_name_label)
        layout.addWidget(full_name_edit)
        layout.addWidget(add_button)
        layout.addWidget(cancel_button)

        add_dialog.setLayout(layout)

        add_button.clicked.connect(lambda: self.add_faculty(
            short_name_edit.text(),
            full_name_edit.text(),
            add_dialog
        ))

        cancel_button.clicked.connect(add_dialog.close)

        add_dialog.exec_()

    def add_faculty(self, short_name, full_name, dialog):
        print(f"Adding faculty: {short_name}")
        try:
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            cursor.execute("""
                INSERT INTO Faculties (facultie_short, facultie_full)
                VALUES (?, ?)
            """, (short_name, full_name))
            conn.commit()
            conn.close()
            self.load_data()
            dialog.close()
        except sqlite3.IntegrityError as e:
            print(f"Add faculty integrity error: {e}")
            QMessageBox.warning(self, "Ошибка", f"Не удалось добавить факультет: {e}")
            conn.close()
        except Exception as e:
            print(f"Add faculty error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка добавления факультета: {e}")
            conn.close()

    def add_division_form(self):
        print("Opening add division form")
        add_dialog = QDialog(self)
        add_dialog.setWindowTitle("Добавить группу")

        division_label = QLabel("Название группы:")
        division_edit = QLineEdit()

        add_button = QPushButton("Добавить")
        cancel_button = QPushButton("Отмена")

        layout = QVBoxLayout()
        layout.addWidget(division_label)
        layout.addWidget(division_edit)
        layout.addWidget(add_button)
        layout.addWidget(cancel_button)

        add_dialog.setLayout(layout)

        add_button.clicked.connect(lambda: self.add_division(division_edit.text(), add_dialog))
        cancel_button.clicked.connect(add_dialog.close)

        add_dialog.exec_()

    def add_division(self, division, dialog):
        print(f"Adding division: {division}")
        try:
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            cursor.execute("""
                INSERT INTO division (division)
                VALUES (?)
            """, (division,))
            conn.commit()
            conn.close()
            self.load_data()
            dialog.close()
        except sqlite3.IntegrityError as e:
            print(f"Add division integrity error: {e}")
            QMessageBox.warning(self, "Ошибка", f"Не удалось добавить группу: {e}")
            conn.close()
        except Exception as e:
            print(f"Add division error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка добавления группы: {e}")
            conn.close()

    def get_divisions(self):
        print("Fetching divisions")
        try:
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT division FROM division WHERE division IS NOT NULL")
            divisions = [row[0] for row in cursor.fetchall()]
            conn.close()
            return divisions
        except Exception as e:
            print(f"Get divisions error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка получения групп: {e}")
            return []

    def get_posts(self):
        print("Fetching posts")
        try:
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT post FROM post WHERE post IS NOT NULL")
            posts = [row[0] for row in cursor.fetchall()]
            conn.close()
            return posts
        except Exception as e:
            print(f"Get posts error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка получения должностей: {e}")
            return []

    def get_faculties(self):
        print("Fetching faculties")
        try:
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()
            cursor.execute("SELECT facultie_short FROM Faculties WHERE facultie_short IS NOT NULL")
            faculties = [row[0] for row in cursor.fetchall()]
            conn.close()
            return faculties
        except Exception as e:
            print(f"Get faculties error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка получения факультетов: {e}")
            return []

    def export_selected_to_excel(self):
        print("Exporting to Excel")
        try:
            selected_items = self.table.selectedItems()
            if not selected_items:
                QMessageBox.warning(self, "Ошибка", "Выберите строки для экспорта")
                return

            file_path, _ = QFileDialog.getSaveFileName(self, "Сохранить как", "", "Excel Files (*.xlsx);;All Files (*)")

            if file_path and not file_path.endswith('.xlsx'):
                file_path += '.xlsx'

            if not file_path:
                return

            selected_rows = set()
            for item in selected_items:
                selected_rows.add(item.row())

            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["ID", "Фамилия", "Имя", "Отчество", "Логин", "Пароль", "Подразделение", "Должность", "Факультет"])

            for row in selected_rows:
                row_data = [self.table.item(row, col).text() if self.table.item(row, col) else '' for col in range(self.table.columnCount())]
                sheet.append(row_data)

            workbook.save(file_path)
        except Exception as e:
            print(f"Export to Excel error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка экспорта в Excel: {e}")

    def export_selected_to_CSV(self):
        print("Exporting to CSV")
        try:
            selected_items = self.table.selectedItems()
            if not selected_items:
                QMessageBox.warning(self, "Ошибка", "Выберите строки для экспорта")
                return

            file_path, _ = QFileDialog.getSaveFileName(self, "Сохранить как", "", "CSV Files (*.csv);;All Files (*)")

            if file_path and not file_path.endswith('.csv'):
                file_path += '.csv'

            if not file_path:
                return

            selected_rows = set()
            for item in selected_items:
                selected_rows.add(item.row())

            with open(file_path, 'w', newline='', encoding='utf-8') as file:
                writer = csv.writer(file)
                writer.writerow(
                    ["ID", "Фамилия", "Имя", "Отчество", "Логин", "Пароль", "Подразделение", "Должность", "Факультет"])

                for row in selected_rows:
                    row_data = [self.table.item(row, col).text() if self.table.item(row, col) else '' for col in range(self.table.columnCount())]
                    writer.writerow(row_data)
        except Exception as e:
            print(f"Export to CSV error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка экспорта в CSV: {e}")

    def import_from_excel(self):
        print("Importing from Excel")
        try:
            file_path, _ = QFileDialog.getOpenFileName(self, "Открыть файл Excel", "",
                                                       "Excel Files (*.xlsx);;All Files (*)")

            if not file_path:
                return

            df = pd.read_excel(file_path)
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()

            required_columns = ['Фамилия', 'Имя', 'Отчество']
            for col in required_columns:
                if col not in df.columns:
                    raise ValueError(f"Отсутствует необходимый столбец: {col}")

            for _, row in df.iterrows():
                last_name = row['Фамилия']
                first_name = row['Имя']
                middle_name = row['Отчество']
                subdivision = row.get('Подразделение', None)
                position = row.get('Должность', None)
                faculty = row.get('Факультет', None)

                login = self.generate_login_import(last_name, first_name, middle_name)
                password = self.generate_password_import()

                cursor.execute("""
                    INSERT INTO user (surname, name, patronymic, login, password, division, post, faculty)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """, (last_name, first_name, middle_name or None, login, password, subdivision or None, position or None, faculty or None))

            conn.commit()
            conn.close()
            self.load_data()
        except Exception as e:
            print(f"Import from Excel error: {e}")
            QMessageBox.critical(self, "Ошибка", f"Ошибка импорта из Excel: {e}")
            conn.close()

    def import_from_csv(self):
        print("Importing from CSV")
        try:
            file_path, _ = QFileDialog.getOpenFileName(self, "Открыть файл CSV", "", "CSV Files (*.csv);;All Files (*)")

            if not file_path:
                return

            df = pd.read_csv(file_path)
            conn = sqlite3.connect(self.current_db_path)
            cursor = conn.cursor()

            required_columns = ['Фамилия', 'Имя', 'Отчество']
            for col in required_columns:
                if col not in df.columns:
                    raise ValueError(f"Отсутствует необходимый столбец: {col}")

            for _, row in df.iterrows():
                last_name = row['Фамилия']
                first_name = row['Имя']
                middle_name = row['Отчество']
                subdivision = row.get('Подразделение', None)
                position = row.get('Должность', None)
                faculty = row.get('Факультет', None)

                login = self.generate_login_import(last_name, first_name, middle_name)
                password = self.generate_password_import()

                cursor.execute("""
                    INSERT INTO user (surname, name, patronymic, login, password, division, post, faculty)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """, (last_name, first_name, middle_name or None, login, password, subdivision or None, position or None, faculty or None))

            conn.commit()
            conn.close()
            self.load_data()
        except Exception as e:
            #print( f"Ошибка импорта из CSV: {e})
            QMessageBox.critical(self, "Ошибка", f"Ошибка импорта: {e}")
            conn.close()

def main():
    print("Starting application")
    try:
        app = QApplication(sys.argv)
        main_window = MainWindow()
        if main_window.current_db_path:
            if main_window.show_login_dialog():
                main_window.show()
                main_window.load_data()
            else:
                print("No login, exiting")
                return
        else:
            main_window.show()
        return app.exec_()
    except Exception as e:
        print(f"Main application error: {e}")
        return 1

if __name__ == "__main__":
    sys.exit(main())